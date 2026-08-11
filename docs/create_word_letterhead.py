import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc_path = r"c:\Users\Kevin\Projects\ML\agriculture\docs\Kuapa_Dwaso_Letterhead_Template.docx"
png_logo_path = r"c:\Users\Kevin\Projects\ML\agriculture\docs\kuapa_dwaso_logo.png"

doc = docx.Document()

# Set standard page margins (0.75 in)
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    # Configure Header
    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Create a 2-column table in the header: Left column for logo + title, right column for contact details
    # Or single block layout matching user screenshot
    
    # Let's add logo image to header
    run_logo = header_para.add_run()
    if os.path.exists(png_logo_path):
        run_logo.add_picture(png_logo_path, width=Inches(0.65))
    
    # Title line
    p_title = header.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("KUAPA DWASO")
    run_title.font.name = "Outfit"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 31, 20) # #0f1f14 ink
    
    # Subtitle line
    p_sub = header.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    run_sub = p_sub.add_run("Student-Led Agritech Innovation Initiative")
    run_sub.font.name = "Plus Jakarta Sans"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(107, 114, 128) # gray 500
    
    # Contact bar
    p_contact = header.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(8)
    
    run_url = p_contact.add_run("kuapadwaso.com")
    run_url.font.name = "Plus Jakarta Sans"
    run_url.font.size = Pt(9.5)
    run_url.font.bold = True
    run_url.font.color.rgb = RGBColor(45, 138, 78) # #2d8a4e field green
    
    run_sep1 = p_contact.add_run("  |  ")
    run_sep1.font.name = "Plus Jakarta Sans"
    run_sep1.font.size = Pt(9.5)
    run_sep1.font.color.rgb = RGBColor(212, 168, 67) # #d4a843 gold
    
    run_email = p_contact.add_run("info@kuapadwaso.com")
    run_email.font.name = "Plus Jakarta Sans"
    run_email.font.size = Pt(9.5)
    run_email.font.bold = True
    run_email.font.color.rgb = RGBColor(15, 31, 20)
    
    run_sep2 = p_contact.add_run("  |  ")
    run_sep2.font.name = "Plus Jakarta Sans"
    run_sep2.font.size = Pt(9.5)
    run_sep2.font.color.rgb = RGBColor(212, 168, 67)
    
    run_phone = p_contact.add_run("054 903 7907")
    run_phone.font.name = "Plus Jakarta Sans"
    run_phone.font.size = Pt(9.5)
    run_phone.font.bold = True
    run_phone.font.color.rgb = RGBColor(15, 31, 20)
    
    # Add bottom border XML line to header paragraph
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="12" w:space="4" w:color="2D8A4E"/>'
                     r'</w:pBdr>')
    p_contact._p.get_or_add_pPr().append(pBdr)

# Add sample body text
p_body = doc.add_paragraph()
p_body.paragraph_format.space_before = Pt(20)
run_body = p_body.add_run("Start typing your document content here...")
run_body.font.name = "Plus Jakarta Sans"
run_body.font.size = Pt(11)

doc.save(doc_path)
print(f"Saved docx template to {doc_path}")
